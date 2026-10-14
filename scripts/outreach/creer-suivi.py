"""Crée le tableur vierge de suivi des contacts Instagram (influenceurs, bars à jeux…).

Sortie : _outreach/KYRAN-suivi-influenceurs.xlsx (dossier ignoré par git : données privées).
Usage : python scripts/outreach/creer-suivi.py [--force]
"""
import datetime, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule, DataBarRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

F="Arial"; N=300
INK="1F2A44"; GOLD="C8962E"; CREAM="FBF6EC"; INPUT="FFF7D6"
wb=Workbook()

# ---- Listes
L=wb.active; L.title="Listes"
lists={"Type":["Influenceur","Bar à jeux","Boutique de jeux","Ludothèque","Média / blog","Association","Autre"],
"Oui/Non":["Oui","Non"],
"Publication":["Oui","Non","Prévue","Refus"],
"Retour":["Très positif","Positif","Neutre","Négatif","Pas encore joué"],
"Réseau":["Instagram","TikTok","YouTube","Facebook","Email","Autre"]}
for c,(k,v) in enumerate(lists.items(),1):
    L.cell(1,c,k).font=Font(name=F,bold=True)
    for r,x in enumerate(v,2): L.cell(r,c,x).font=Font(name=F)
    L.column_dimensions[chr(64+c)].width=18

# ---- Suivi DM
S=wb.create_sheet("Suivi DM",0)
cols=[("N°",6),("Statut (auto)",20),("Nom / pseudo",22),("Type",16),("Compte Instagram",20),("URL Instagram",34),
("Abonnés Insta",13),("Ville",14),("Date 1er DM",13),("A répondu ?",11),("Date réponse",13),("Résumé de sa réponse",40),
("Jeu envoyé ?",11),("Nb jeux envoyés",12),("Date d'envoi",13),("Retour sur le jeu",16),("Publication faite ?",13),
("Date publication",13),("Lien publication",34),("Site web / Google Maps (bar)",34),("Email / contact",26),
("Prochaine action",28),("Notes",36)]
S.row_dimensions[1].height=34; S.row_dimensions[2].height=40
S["A1"]="KYRAN — Suivi des contacts Instagram (@kyran.jeu)"
S["A1"].font=Font(name=F,size=16,bold=True,color=INK)
S["F1"]="Cases jaunes = à remplir · « Statut » se calcule tout seul · ligne 3 = EXEMPLE à effacer · coller les données de Claude à partir de la colonne C"
S["F1"].font=Font(name=F,size=9,italic=True,color="6B6B6B")
thin=Side(style="thin",color="D9CFBD")
for i,(h,w) in enumerate(cols,1):
    c=S.cell(2,i,h); c.font=Font(name=F,bold=True,color="FFFFFF",size=10)
    c.fill=PatternFill("solid",fgColor=INK); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    S.column_dimensions[c.column_letter].width=w
last=2+N
for r in range(3,last+1):
    S.cell(r,1,f'=IF(C{r}="","",ROW()-2)')
    S.cell(r,2,f'=IF(C{r}="","",IF(M{r}="Oui",IF(Q{r}="Oui","Jeu envoyé + publié","Jeu envoyé"),IF(J{r}="Oui","Réponse, pas d\'envoi","DM sans réponse")))')
    for c in range(1,len(cols)+1):
        cell=S.cell(r,c); cell.font=Font(name=F,size=10); cell.border=Border(bottom=thin)
        cell.alignment=Alignment(vertical="top",wrap_text=c in (12,22,23))
    for c in (9,11,15,18): S.cell(r,c).number_format="dd/mm/yyyy"
    S.cell(r,7).number_format='#,##0'; S.cell(r,14).number_format='0'
# example row
ex=["", "", "Exemple Ludique", "Influenceur","@exemple.ludique","https://www.instagram.com/exemple.ludique/",12500,"Lyon",
"2026-10-02","Oui","2026-10-03","Partante pour tester, demande 2 boîtes pour un concours","Oui",2,"2026-10-05","Très positif",
"Prévue","","","","exemple@mail.fr","Relancer le 20/10 pour la date du post","EXEMPLE — supprimer cette ligne"]
for c,v in enumerate(ex,1):
    if c in (1,2) or v=="": continue
    if c in (9,11,15): v=datetime.date.fromisoformat(v)
    S.cell(3,c,v)
for c in range(1,len(cols)+1): S.cell(3,c).font=Font(name=F,size=10,italic=True,color="8A8A8A")
S.cell(3,6).hyperlink=None
# input fill (yellow light) on editable columns
for r in range(3,last+1):
    for c in range(1,len(cols)+1):
        if c not in (1,2): S.cell(r,c).fill=PatternFill("solid",fgColor=INPUT)
S.freeze_panes="D3"
S.auto_filter.ref=f"A2:{S.cell(2,len(cols)).column_letter}{last}"
def dv(formula,col):
    d=DataValidation(type="list",formula1=formula,allow_blank=True); S.add_data_validation(d); d.add(f"{col}3:{col}{last}")
dv("=Listes!$A$2:$A$8","D"); dv("=Listes!$B$2:$B$3","J"); dv("=Listes!$B$2:$B$3","M")
dv("=Listes!$C$2:$C$5","Q"); dv("=Listes!$D$2:$D$6","P")
rng=f"B3:B{last}"
for txt,color in [("Jeu envoyé + publié","2E7D32"),("Jeu envoyé","4CAF50"),("Réponse, pas d'envoi","F0A030"),("DM sans réponse","B0B0B0")]:
    S.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=[f'"{txt.replace(chr(39),chr(39))}"'],
        fill=PatternFill("solid",fgColor=color),font=Font(name=F,bold=True,color="FFFFFF")))
S.conditional_formatting.add(f"G3:G{last}",DataBarRule(start_type="num",start_value=0,end_type="max",color=GOLD))
S["B2"].comment=Comment("Calculé : Jeu envoyé (+ publié) > Réponse, pas d'envoi > DM sans réponse. Ne pas saisir.","Claude")

# ---- Tableau de bord
D=wb.create_sheet("Tableau de bord",0)
D.sheet_view.showGridLines=False
D.column_dimensions["A"].width=3; D.column_dimensions["B"].width=34; D.column_dimensions["C"].width=16
D.column_dimensions["D"].width=3; D.column_dimensions["E"].width=26
for k in "FGHI": D.column_dimensions[k].width=14
D["B2"]="KYRAN — Tableau de bord influence"; D["B2"].font=Font(name=F,size=18,bold=True,color=INK)
D["B3"]="Tout se met à jour depuis l'onglet « Suivi DM »"; D["B3"].font=Font(name=F,size=9,italic=True,color="6B6B6B")
R="'Suivi DM'!"
def rg(c): return f"{R}${c}$3:${c}${last}"
kpi=[("Personnes contactées",f'=COUNTA({rg("C")})',"0"),
("Ont répondu",f'=COUNTIF({rg("J")},"Oui")',"0"),
("Taux de réponse",f'=IF(C5=0,0,C6/C5)',"0%"),
("DM sans réponse",f'=COUNTIF({rg("B")},"DM sans réponse")',"0"),
("Réponse, pas d'envoi",f'=COUNTIF({rg("B")},"Réponse, pas d\'envoi")',"0"),
("Personnes ayant reçu un jeu",f'=COUNTIF({rg("M")},"Oui")',"0"),
("Nombre total de jeux envoyés",f'=SUM({rg("N")})',"0"),
("Publications faites",f'=COUNTIF({rg("Q")},"Oui")',"0"),
("Publications prévues",f'=COUNTIF({rg("Q")},"Prévue")',"0"),
("Taux de publication (sur jeux envoyés)",f'=IF(C10=0,0,C12/C10)',"0%"),
("Audience cumulée touchée (abonnés)",f'=SUMIFS({rg("G")},{rg("Q")},"Oui")',"#,##0"),
("Coût estimé des jeux envoyés (€)",f'=C11*C19+C10*C20','#,##0.00 €')]
D["B4"]="Indicateur"; D["C4"]="Valeur"
for c in ("B4","C4"):
    D[c].font=Font(name=F,bold=True,color="FFFFFF"); D[c].fill=PatternFill("solid",fgColor=INK)
for i,(lab,f,fmt) in enumerate(kpi,5):
    D.cell(i,2,lab).font=Font(name=F,size=11)
    c=D.cell(i,3,f); c.font=Font(name=F,size=12,bold=True,color=INK); c.number_format=fmt
    c.alignment=Alignment(horizontal="right")
    fill=PatternFill("solid",fgColor=CREAM if i%2 else "FFFFFF")
    D.cell(i,2).fill=fill; c.fill=fill
D["B18"]="Hypothèses de coût"; D["B18"].font=Font(name=F,bold=True,color=INK)
D["B19"]="Prix d'une boîte (€)"; D["C19"]=9.99
D["B20"]="Frais de port par colis (€)"; D["C20"]=2.99
for c in ("C19","C20"):
    D[c].font=Font(name=F,color="0000FF"); D[c].fill=PatternFill("solid",fgColor="FFFF00"); D[c].number_format='0.00 €'
D["B19"].font=D["B20"].font=Font(name=F)
D["C19"].comment=Comment("Prix public KYRAN (9,99 €) et port (2,99 €) du lien Stripe, d'après CLAUDE.md. Modifiable.","Claude")
D["B21"]="Hypothèse : un colis par personne ; prix public, pas le coût de fabrication."
D["B21"].font=Font(name=F,size=9,italic=True,color="6B6B6B")
# by type
D["E4"]="Par type"
heads=["Contactés","Réponses","Jeux envoyés","Publications"]
for j,h in enumerate(["Par type"]+heads):
    c=D.cell(4,5+j,h); c.font=Font(name=F,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=GOLD)
    c.alignment=Alignment(horizontal="center")
for i,t in enumerate(lists["Type"],5):
    D.cell(i,5,t).font=Font(name=F)
    fs=[f'=COUNTIF({rg("D")},E{i})',f'=COUNTIFS({rg("D")},E{i},{rg("J")},"Oui")',
        f'=SUMIFS({rg("N")},{rg("D")},E{i})',f'=COUNTIFS({rg("D")},E{i},{rg("Q")},"Oui")']
    for j,f in enumerate(fs): 
        c=D.cell(i,6+j,f); c.font=Font(name=F); c.alignment=Alignment(horizontal="center")
tr=5+len(lists["Type"])
D.cell(tr,5,"Total").font=Font(name=F,bold=True)
for j in range(4):
    col=chr(70+j); c=D.cell(tr,6+j,f"=SUM({col}5:{col}{tr-1})"); c.font=Font(name=F,bold=True); c.alignment=Alignment(horizontal="center")
    c.border=Border(top=Side(style="thin",color=INK))
D.cell(tr+2,5,"Note : la ligne EXEMPLE de « Suivi DM » compte tant qu'elle n'est pas effacée.").font=Font(name=F,size=9,italic=True,color="AA3333")

# ---- Mode d'emploi
M=wb.create_sheet("Mode d'emploi")
M.column_dimensions["A"].width=110; M.sheet_view.showGridLines=False
txt=["Mode d'emploi",
"1. Une ligne par personne ou lieu contacté (influenceur, bar à jeux, boutique…).",
"2. Remplir les cases jaunes (ou coller la sortie de Claude for Chrome en C3). Les listes déroulantes évitent les fautes (Type, Oui/Non, Publication, Retour).",
"3. « Statut (auto) » se calcule : Jeu envoyé + publié / Jeu envoyé / Réponse, pas d'envoi / DM sans réponse.",
"4. URL Instagram : https://www.instagram.com/<pseudo>/ · Bars : mettre le site ou la fiche Google Maps.",
"5. Dates au format JJ/MM/AAAA. Abonnés en nombre (12500, pas « 12,5 k »).",
"6. L'onglet « Tableau de bord » se met à jour seul. Les hypothèses de coût (cases jaunes) sont modifiables.",
"7. Effacer la ligne EXEMPLE (ligne 3) avant usage réel.",
"8. Remplissage automatique : python scripts/outreach/remplir-suivi.py contacts.json (voir docs/SUIVI-INSTAGRAM.md)."]
for i,t in enumerate(txt,1):
    M.cell(i,1,t).font=Font(name=F,size=14 if i==1 else 11,bold=i==1,color=INK if i==1 else "000000")
wb.move_sheet("Listes",offset=10)
out=Path(__file__).resolve().parents[2]/"_outreach"/"KYRAN-suivi-influenceurs.xlsx"
if out.exists() and "--force" not in sys.argv:
    sys.exit(f"{out} existe déjà (ajouter --force pour l'écraser).")
out.parent.mkdir(exist_ok=True)
wb.save(out); print(out)
