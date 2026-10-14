"""Remplit l'onglet « Suivi DM » du tableur à partir d'un fichier JSON de contacts.

Usage : python scripts/outreach/remplir-suivi.py contacts.json [--ecraser]

contacts.json : liste d'objets dont les clés sont celles de CHAMPS ci-dessous (toutes facultatives
sauf « nom » ou « compte »). Un contact déjà présent (même compte Instagram, sinon même nom) est
complété : seules ses cases vides sont remplies, sauf avec --ecraser. Les nouveaux contacts sont
ajoutés à la suite. La ligne EXEMPLE du modèle est effacée au premier remplissage.
"""
import datetime, json, re, sys
from pathlib import Path
from openpyxl import load_workbook

FICHIER = Path(__file__).resolve().parents[2] / "_outreach" / "KYRAN-suivi-influenceurs.xlsx"
PREMIERE, DERNIERE = 3, 302

# clé JSON -> colonne de « Suivi DM » (A = N° et B = Statut sont des formules)
CHAMPS = {
    "nom": "C", "type": "D", "compte": "E", "url_instagram": "F", "abonnes": "G", "ville": "H",
    "date_premier_dm": "I", "a_repondu": "J", "date_reponse": "K", "resume_reponse": "L",
    "jeu_envoye": "M", "nb_jeux": "N", "date_envoi": "O", "retour_jeu": "P", "publication": "Q",
    "date_publication": "R", "lien_publication": "S", "site_web": "T", "contact": "U",
    "prochaine_action": "V", "notes": "W",
}
DATES = {"date_premier_dm", "date_reponse", "date_envoi", "date_publication"}
NOMBRES = {"abonnes", "nb_jeux"}


def cle(compte, nom):
    compte = (compte or "").strip().lstrip("@").lower()
    return ("compte", compte) if compte else ("nom", (nom or "").strip().lower())


def convertir(champ, valeur):
    if valeur in (None, ""):
        return None
    if champ in DATES and isinstance(valeur, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.datetime.strptime(valeur.strip(), fmt).date()
            except ValueError:
                pass
        return valeur
    if champ in NOMBRES and isinstance(valeur, str):
        texte = valeur.strip().lower().replace(" ", "").replace(" ", "").replace(",", ".")
        m = re.fullmatch(r"([\d.]+)(k|m)?", texte)
        if m:
            n = float(m.group(1)) * {"k": 1_000, "m": 1_000_000}.get(m.group(2), 1)
            return int(round(n))
        return valeur
    if champ == "compte" and isinstance(valeur, str) and not valeur.startswith("@"):
        return "@" + valeur.strip()
    return valeur


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    ecraser = "--ecraser" in sys.argv
    contacts = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    if not FICHIER.exists():
        sys.exit(f"{FICHIER} introuvable : lancer d'abord scripts/outreach/creer-suivi.py")

    wb = load_workbook(FICHIER)
    ws = wb["Suivi DM"]
    if ws["C3"].value == "Exemple Ludique":
        for col in CHAMPS.values():
            ws[f"{col}3"].value = None

    index, libre = {}, []
    for r in range(PREMIERE, DERNIERE + 1):
        if ws[f"C{r}"].value or ws[f"E{r}"].value:
            index[cle(ws[f"E{r}"].value, ws[f"C{r}"].value)] = r
        else:
            libre.append(r)

    ajoutes = completes = 0
    for c in contacts:
        k = cle(c.get("compte"), c.get("nom"))
        if not k[1]:
            continue
        if k in index:
            r = index[k]
            completes += 1
        else:
            if not libre:
                sys.exit("Plus de lignes libres (300 max) : agrandir le modèle.")
            r = libre.pop(0)
            index[k] = r
            ajoutes += 1
        for champ, col in CHAMPS.items():
            v = convertir(champ, c.get(champ))
            if v is None:
                continue
            cellule = ws[f"{col}{r}"]
            if cellule.value in (None, "") or ecraser:
                cellule.value = v
                if isinstance(v, str) and v.startswith("http"):
                    cellule.hyperlink = v

    wb.save(FICHIER)
    print(f"{ajoutes} contact(s) ajouté(s), {completes} complété(s) → {FICHIER}")


if __name__ == "__main__":
    main()
